"""Public read verification; auth and rejected archive-resume probes never generate content."""
import argparse
import hashlib
import io
import json
from pathlib import Path
import re
import time
import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs/自然群聊-20261015'
BASE = 'https://vimax.217.194.133.194.sslip.io'


def fingerprint(messages):
    return hashlib.sha256(json.dumps(messages, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def session_labels(client):
    response = client.get('/api/sessions')
    response.raise_for_status()
    values = response.json()
    if not isinstance(values, list):
        values = values.get('sessions', [])
    if isinstance(values, dict):
        values = list(values.values())
    result = {}
    for item in values:
        sid = item.get('sessionId') or item.get('session_id') or item.get('id') or ''
        if sid.startswith('dialogue-'):
            result[sid.removeprefix('dialogue-')] = item.get('projectName') or item.get('project_name')
    return result


def get_run(client, run_id):
    response = client.get('/dialogue-api/simulations/' + run_id, params={'view': 'workspace'})
    response.raise_for_status()
    run = response.json()
    assert run['id'] == run_id
    return run


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('mode', choices=('before', 'after'))
    parser.add_argument('--selected', default='V40-R001')
    parser.add_argument('--release')
    args = parser.parse_args()
    baseline_path = OUT / '公网发布前历史.json'
    key_file = ROOT / '.vimax/deploy/登录信息.txt'
    key = re.search(r'^访问密钥[:：]\s*(.+)$', key_file.read_text(encoding='utf-8'), re.M).group(1).strip()
    proof = {'url': BASE, 'mode': args.mode, 'release': args.release, 'records': [], 'checks': {}}
    with httpx.Client(base_url=BASE, timeout=40, trust_env=False, follow_redirects=True, headers={'Origin': BASE}) as client:
        assert client.get('/api/sessions').status_code == 401
        client.post('/auth/login', json={'accessKey': key}).raise_for_status()
        try:
            labels = session_labels(client)
            if args.mode == 'after':
                baseline = json.loads(baseline_path.read_text(encoding='utf-8'))
                originals = {r['id']: r for r in baseline['records']}
                entries = json.loads((OUT / 'versions.json').read_text(encoding='utf-8'))
                new = {v['run_id']: v for v in entries if int(v['version'][1:3]) >= 4}
                assert set(labels) == set(originals) | set(new), 'Visible history set changed unexpectedly'
                page = client.get('/')
                page.raise_for_status()
                assert 'www-authenticate' not in page.headers
                assets = re.findall(r'(?:src|href)="(/assets/[^\"]+)"', page.text)
                assert assets
                for asset in assets:
                    remote = client.get(asset)
                    remote.raise_for_status()
                    assert remote.content == (ROOT / 'web/dist' / asset.lstrip('/')).read_bytes()
                proof['checks']['frontend_assets_exact_match'] = assets
            for run_id in sorted(labels):
                started = time.monotonic()
                run = get_run(client, run_id)
                hashed = fingerprint(run['messages'])
                record = {'id': run_id, 'label': labels[run_id], 'status': run['status'],
                          'messages': len(run['messages']), 'sha256': hashed,
                          'workspace_seconds': round(time.monotonic()-started, 3)}
                if args.mode == 'after':
                    if run_id in originals:
                        old = originals[run_id]
                        assert all(record[k] == old[k] for k in ('id', 'label', 'status', 'messages', 'sha256'))
                    else:
                        entry = new[run_id]
                        folder = OUT / entry['version']
                        local = json.loads((folder / 'run.json').read_text(encoding='utf-8'))
                        for field in ('messages', 'status', 'config', 'errors', 'version_record', 'program_review_deferred'):
                            assert run.get(field) == local.get(field), (entry['version'], field)
                        publication = run['publication']
                        assert publication['read_only'] and publication['version'] == entry['version']
                        assert publication['source_run_sha256'] == hashlib.sha256((folder / 'run.json').read_bytes()).hexdigest()
                        assert publication['human_verdict'] == entry['human_verdict']
                        assert publication['program_review_state'] == ('deferred' if local.get('program_review_deferred') else 'reviewed' if local.get('program_review_history') else 'not_completed')
                        assert run['template']['source'] == json.loads((folder / 'template.json').read_text(encoding='utf-8'))['source']
                        record['version'] = entry['version']
                        record['program_review_state'] = publication['program_review_state']
                proof['records'].append(record)
            if args.mode == 'after':
                selected = next(v for v in new.values() if v['version'] == args.selected)
                selected_run = get_run(client, selected['run_id'])
                assert selected_run['status'] == 'completed' and selected_run['publication']['selected']
                response = client.get('/dialogue-api/simulations/' + selected['run_id'] + '/export/xlsx')
                response.raise_for_status()
                actual = load_workbook(io.BytesIO(response.content)).active
                expected = load_workbook(OUT / args.selected / '群聊150条.xlsx').active
                assert (actual.max_row, actual.max_column) == (151, 9)
                assert list(actual.values) == list(expected.values)
                proof['checks']['xlsx'] = '151x9; all cells match local'
                for version in ('V06', 'V40'):
                    entry = next(v for v in new.values() if v['version'] == version)
                    before = get_run(client, entry['run_id'])
                    rejected = client.post('/dialogue-api/simulations/' + entry['run_id'] + '/resume')
                    assert rejected.status_code == 409 and '已归档版本' in rejected.json()['detail']
                    assert get_run(client, entry['run_id']) == before
                proof['checks'].update(original_records_unchanged=len(originals), new_records_exact_match=len(new),
                                       archived_resume_rejected_without_mutation=True)
            health = client.get('/dialogue-api/health')
            health.raise_for_status()
            proof['checks']['health'] = health.json()
        finally:
            client.post('/auth/logout').raise_for_status()
        assert client.get('/api/sessions').status_code == 401
    proof['checks']['key_login_logout'] = True
    destination = baseline_path if args.mode == 'before' else OUT / '公网最终验证.json'
    assert not destination.exists(), 'Keep existing verification evidence immutable'
    destination.write_text(json.dumps(proof, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'mode': args.mode, 'records': len(proof['records']), 'checks': proof['checks']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
