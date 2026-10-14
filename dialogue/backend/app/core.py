from __future__ import annotations

import io
import re
from collections import Counter
from difflib import SequenceMatcher
from typing import Any

from openpyxl import Workbook, load_workbook

from .locale import message_char_limit
from .dialogue_quality import closes_group

HEADERS = ["顺序", "发送人", "回复", "消息类型", "延时", "内容", "图片码", "图片方式", "替换日期"]
CODE_RE = re.compile(r"【([A-Za-z][A-Za-z0-9_-]{2,})】|(?<!\d)\d{5,}(?!\d)")
URL_RE = re.compile(r"https?://[^\s,，。、;；<>】\"'“”‘’（）()「」『』]+(?<![.!?])")
BRACKET_RE = re.compile(r"【[^】]{2,80}】")
EMAIL_RE = re.compile(r"(?<![\w.+-])[\w.+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
META_RE = re.compile(r"本地活动交流群|本地模拟|模拟群聊|生成剧本|剧本对话|母本|原表|原文|原说明|现有说明|我是AI|AI模型|Agent", re.IGNORECASE)
CODE_PAGE_RE = re.compile(r"(?:页面|入口|任务页)[^，。！？\n]{0,14}(?:显示|生成|提供|复制|找到|取得|有|没显示)[^，。！？\n]{0,12}(?:原代码|代码)|(?:原代码|代码)[^，。！？\n]{0,12}(?:从|在|按)[^，。！？\n]{0,10}(?:页面|入口|任务页)")
GROUP_SCREENSHOT_DENIAL_RE = re.compile(
    r"(?:不需要|无需|不用|不必|不要)[^。！？；\n]{0,18}(?:群里|群组|群内|工作组)[^。！？；\n]{0,15}(?:截图|截屏|图片)"
    r"|(?:截图|截屏|图片)[^。！？；\n]{0,18}(?:不需要|无需|不用|不必|不要)[^。！？；\n]{0,15}(?:群里|群组|群内|工作组)"
)
EARLY_ACTION_RE = re.compile(r"(?:我|这边)[^。！？\n]{0,16}(?:已经|已|刚|正在|现在按说明)[^。！？\n]{0,12}(?:打开|进入|点赞|截[图屏]|提交)")
CONTACT_ACCEPTANCE_WAIT_RE = re.compile(
    r"(?:等|待).{0,18}(?:通过|同意|接受|确认).{0,60}(?:发|提交)"
    r"|(?:通过|确认)(?:前|后).{0,60}(?:发|提交)"
)
OPERATOR_APPROVAL_WAIT_RE = re.compile(
    r"(?:等|待)坐席.{0,8}(?:确认|许可|允许)"
)


def read_source(data: bytes) -> dict[str, Any]:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 XLSX 文件") from exc
    if not book.worksheets:
        raise ValueError("工作簿没有工作表")
    sheet = book.worksheets[0]
    headers = [sheet.cell(1, col).value for col in range(1, 10)]
    if headers != HEADERS:
        raise ValueError(f"第一张工作表需要九列表头：{'、'.join(HEADERS)}")
    messages: list[dict[str, Any]] = []
    seen: set[int] = set()
    for source_row, values in enumerate(sheet.iter_rows(min_row=2, max_col=9, values_only=True), 2):
        if all(value is None for value in values):
            continue
        try:
            seq = int(values[0])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"第 {source_row} 行顺序不是整数") from exc
        if seq in seen or seq <= 0:
            raise ValueError(f"第 {source_row} 行顺序重复或无效")
        seen.add(seq)
        if not values[1] or values[3] not in {"文本", "图片", "图片+文本"}:
            raise ValueError(f"第 {source_row} 行发送人或消息类型无效")
        reply = None
        if values[2] is not None and str(values[2]).strip():
            try:
                reply = int(values[2])
            except (TypeError, ValueError) as exc:
                raise ValueError(f"第 {source_row} 行回复编号无效") from exc
            if reply >= seq:
                raise ValueError(f"第 {source_row} 行回复必须指向之前的消息")
        if "图片" in values[3] and not values[6]:
            raise ValueError(f"第 {source_row} 行图片消息缺少图片码")
        messages.append({
            "seq": seq, "source_row": source_row, "sender": str(values[1]),
            "reply_to": reply, "type": values[3], "delay": str(values[4]) if values[4] is not None else None,
            "content": str(values[5]) if values[5] is not None else None,
            "image_code": str(values[6]) if values[6] is not None else None,
            "image_mode": str(values[7]) if values[7] is not None else None,
            "replace_date": str(values[8]) if values[8] is not None else None,
        })
    messages.sort(key=lambda item: item["seq"])
    ids = {item["seq"] for item in messages}
    if any(item["reply_to"] is not None and item["reply_to"] not in ids for item in messages):
        raise ValueError("存在指向不存在消息的回复编号")
    if not messages or not any(item["sender"] == "坐席" for item in messages):
        raise ValueError("母本必须包含消息和坐席行")
    return {"sheet": sheet.title, "messages": messages, "accounts": sorted({m["sender"] for m in messages if m["sender"] != "坐席"}),
            "codes": sorted({match.group(1) or match.group(0) for m in messages for match in CODE_RE.finditer(m["content"] or "")})}


def analyze(source: dict[str, Any]) -> dict[str, Any]:
    seat = [m for m in source["messages"] if m["sender"] == "坐席"]
    action_terms = ("步骤", "打开链接", "访问链接", "截图", "联系管理员", "添加", "代码", "提交")
    evidence = [m["source_row"] for m in seat if any(term in (m["content"] or "") for term in action_terms)]
    reward = any("奖励" in (m["content"] or "") for m in seat)
    interaction = any("点赞" in (m["content"] or "") for m in seat)
    primary = "引导参与者完成商品互动、提交截图并联系指定人员" if interaction else "引导参与者按母本步骤完成目标动作"
    if reward:
        primary += "以申请奖励"
    return {"goal_type": "引导操作" if evidence else "待确认", "primary_goal": primary if evidence else "请根据母本确认主要目标",
            "secondary_goals": ["解释流程并回答疑问"],
            "completion_actions": ["完成母本指定操作", "提交证明材料", "联系指定人员"] if evidence else [],
            "workflow_steps": ["查看活动入口和参与要求", "完成指定商品互动", "保留完成截图", "联系指定人员并附原代码提交", "等待核验结果"] if evidence else [],
            "evidence_rows": evidence[:12], "confirmed": False,
            "note": "目标来自母本措辞；奖励和付款陈述不代表已核实的外部事实。"}


def min_total(source: dict[str, Any], customer_count: int) -> int:
    return 3  # Seat opening, a participant response, and seat closure.


def build_schedule(source: dict[str, Any], customer_count: int, target: int) -> list[dict[str, Any]]:
    minimum = min_total(source, customer_count)
    if target < minimum:
        raise ValueError(f"至少需要 {minimum} 条消息，才能保留坐席开场、参与者回应和收束")
    return [{"kind": "seat" if index in {0, target - 1} else "open",
             "sender": "坐席" if index in {0, target - 1} else None,
             "stage": "自然群聊", "source_seq": None, "required_literals": []}
            for index in range(target)]


def source_literals(source: dict[str, Any]) -> list[str]:
    seat_text = "\n".join(m["content"] or "" for m in source["messages"] if m["sender"] == "坐席")
    return list(dict.fromkeys(URL_RE.findall(seat_text) + BRACKET_RE.findall(seat_text)))


def protected_values(text: str) -> Counter[str]:
    """Exact values a content-only revision must leave in place."""
    values = URL_RE.findall(text) + BRACKET_RE.findall(text) + EMAIL_RE.findall(text)
    values += [match.group(1) or match.group(0) for match in CODE_RE.finditer(text)]
    return Counter(values)


def unannounced_literals(source: dict[str, Any], text: str, history: list[dict]) -> list[str]:
    """Participants may only repeat protected values already published by the host."""
    announced = "\n".join(str(item.get("content") or "") for item in history if item.get("sender") == "坐席")
    return [value for value in source_literals(source) if value in text and value not in announced]


def group_screenshot_required(source: dict[str, Any]) -> bool:
    return any(m["sender"] == "坐席" and any(word in (m["content"] or "") for word in ("截图", "截屏", "捕获", "捕捉"))
               and any(word in (m["content"] or "") for word in ("群里", "群组", "群内", "工作组"))
               and any(word in (m["content"] or "") for word in ("发", "发布", "提交", "分享"))
               for m in source["messages"])


def contradicts_group_screenshot(source: dict[str, Any], text: str) -> bool:
    return group_screenshot_required(source) and bool(GROUP_SCREENSHOT_DENIAL_RE.search(text))


def claims_action_before_entry(text: str) -> bool:
    return bool(EARLY_ACTION_RE.search(text))


def codes_in(text: str | None) -> set[str]:
    return {match.group(1) or match.group(0) for match in CODE_RE.finditer(text or "")}


def unsupported_code_origin(source: dict[str, Any], text: str) -> bool:
    if not source["codes"] or not CODE_PAGE_RE.search(text):
        return False
    if re.search(r"(?:不要|不用|别|不是)[^，。！？\n]{0,8}(?:从|在|按)?页面", text):
        return False
    return not any(CODE_PAGE_RE.search(m["content"] or "") for m in source["messages"] if m["sender"] == "坐席")


def unsupported_contact_acceptance_wait(source: dict[str, Any], text: str) -> bool:
    """Do not invent contact approval as a prerequisite for sending materials."""
    if (OPERATOR_APPROVAL_WAIT_RE.search(text)
            and any(word in text for word in ("截图", "代码", "材料", "经理", "管理员", "私聊"))):
        if (re.search(r"我(?:想|决定|先)?(?:等|等待)坐席", text)
                or "?" in text or "？" in text
                or re.search(r"(?:不必|不用|无需|别|不需要).{0,8}等", text)):
            return False
        return not any(OPERATOR_APPROVAL_WAIT_RE.search(m["content"] or "")
                       for m in source["messages"] if m["sender"] == "坐席")
    if not CONTACT_ACCEPTANCE_WAIT_RE.search(text) or not any(
            word in text for word in ("截图", "截屏", "代码", "材料")):
        return False
    if "?" in text or "？" in text or re.search(r"(?:不必|不用|无需|别|不需要).{0,8}等", text):
        return False
    return not any(CONTACT_ACCEPTANCE_WAIT_RE.search(m["content"] or "") and any(
        word in (m["content"] or "") for word in ("截图", "截屏", "代码", "材料"))
        for m in source["messages"] if m["sender"] == "坐席")


def unsupported_global_verification_halt(source: dict[str, Any], text: str) -> bool:
    """Do not add an external verification gate to the host's published workflow."""
    verification = re.compile(r"核实|核验|查证|独立确认|独立验证")
    halt = re.compile(r"(?:别|不要|不能|不用|禁止|暂停|停止).{0,12}(?:点击|点开|打开|操作|点赞|截图|提交|发送|代码)|"
                      r"(?:只能|先停在|先停留在).{0,18}(?:确定|第一步|了解|讨论)")
    if not verification.search(text) or not halt.search(text):
        return False
    if re.search(r"(?:不必|不用|无需|不需要)\s*(?:再|继续)?(?:等|等待)[^。！？；;]{0,8}(?:核实|核验|查证|独立确认|独立验证)", text):
        return False
    return not any(verification.search(str(item.get("content") or ""))
                   and halt.search(str(item.get("content") or ""))
                   for item in source.get("messages", []) if item.get("sender") == "坐席")


def unsupported_page_detail(source: dict[str, Any], text: str) -> bool:
    """Reject invented UI observations while allowing the source's basic actions."""
    source_text = "\n".join(str(item.get("content") or "") for item in source.get("messages", []))
    details = (r"横幅|弹窗|按钮.{0,10}(?:变成|变为|显示|灰色|绿色)|商品图和按钮",
               r"banner|pop[- ]?up|button.{0,16}(?:turn(?:ed|s)?|chang(?:ed|es)|green|grey|gray)",
               r"バナー|ポップアップ|ボタン.{0,10}(?:変わ|緑|灰色)")
    return any(re.search(pattern, text, re.I) and not re.search(pattern, source_text, re.I)
               for pattern in details)


def validate_run(source: dict[str, Any], run: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    if any("stage" not in slot for slot in run.get("schedule", [])):
        return ["这份结果使用旧版逐行排程，逻辑校验不适用；请重试生成新版本"]
    emitted = run["messages"]
    present = {m["seq"] for m in emitted}
    allowed_literals = set(source_literals(source))
    source_has_entry = any(URL_RE.search(m["content"] or "") for m in source["messages"] if m["sender"] == "坐席")
    entry_provided = False
    for item in emitted:
        if run.get("architecture") == "agent_team" and len(item.get("content") or "") > message_char_limit(
                run.get("config", {}).get("language"), seat=item["sender"] == "坐席"):
            errors.append(f"消息 {item['seq']} 超过群聊单条长度")
        if item.get("reply_to") is not None and (item["reply_to"] not in present or item["reply_to"] >= item["seq"]):
            errors.append(f"消息 {item['seq']} 的回复编号无效")
        if not codes_in(item.get("content")).issubset(set(source["codes"])):
            errors.append(f"消息 {item['seq']} 出现未知代码")
        text = item.get("content") or ""
        prior = [message for message in emitted if message["seq"] < item["seq"]]
        if item["sender"] != "坐席" and unannounced_literals(source, text, prior):
            errors.append(f"消息 {item['seq']} 在坐席公布前使用了活动链接、联系人或代码")
        if (item["sender"] == "坐席" and item["seq"] < run.get("config", {}).get("target_messages", len(run.get("schedule", [])))
                and run.get("status") != "early_end" and closes_group(text, run.get("config", {}).get("language"))):
            errors.append(f"消息 {item['seq']} 提前结束全群讨论")
        if META_RE.search(text):
            errors.append(f"消息 {item['seq']} 出现制作过程用语")
        if unsupported_code_origin(source, text):
            errors.append(f"消息 {item['seq']} 臆测固定代码来自页面")
        if unsupported_page_detail(source, text):
            errors.append(f"消息 {item['seq']} 编造了母本没有的页面细节")
        if unsupported_contact_acceptance_wait(source, text):
            errors.append(f"消息 {item['seq']} 擅自增加联系或发材料前的确认许可要求")
        if contradicts_group_screenshot(source, text):
            errors.append(f"消息 {item['seq']} 否定了母本要求的群组截图提交步骤")
        if item["sender"] == "坐席" and unsupported_global_verification_halt(source, text):
            errors.append(f"消息 {item['seq']} 把母本没有的外部核实条件加成全群操作前提")
        if source_has_entry and item["sender"] != "坐席" and not entry_provided and claims_action_before_entry(text):
            errors.append(f"消息 {item['seq']} 在坐席提供活动入口前声称已操作")
        if item["sender"] == "坐席" and URL_RE.search(text):
            entry_provided = True
        if set(URL_RE.findall(text) + BRACKET_RE.findall(text)) - allowed_literals:
            errors.append(f"消息 {item['seq']} 出现母本之外的链接或括号代码")
    if run["status"] == "completed":
        if len(emitted) != run["config"]["target_messages"]:
            errors.append("最终消息条数与设定不符")
        actual = Counter(m["sender"] for m in emitted)
        for sender in ["坐席"] + source["accounts"] + [f"客户{i + 1}" for i in range(run["config"]["customer_count"])]:
            if sender == "坐席" and not actual[sender]:
                errors.append(f"{sender} 没有发言")
        seat_text = "\n".join(m.get("content") or "" for m in emitted if m["sender"] == "坐席")
        for literal in allowed_literals:
            if literal not in seat_text:
                errors.append(f"坐席发言遗漏母本链接或代码：{literal}")
        if [m.get("stage") for m in emitted] != [slot["stage"] for slot in run["schedule"]]:
            errors.append("对话阶段顺序与计划不符")
        for index, item in enumerate(emitted):
            text = (item.get("content") or "").strip()
            if len(text) < 24:
                continue
            for previous in emitted[max(0, index - 10):index]:
                other = (previous.get("content") or "").strip()
                if len(other) >= 24 and SequenceMatcher(None, text, other).ratio() > 0.88:
                    errors.append(f"消息 {item['seq']} 与消息 {previous['seq']} 内容高度重复")
                    break
    return errors


def export_xlsx(run: dict[str, Any]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "生成剧本"
    sheet.append(HEADERS)
    for m in run["messages"]:
        sheet.append([m["seq"], m["sender"], m.get("reply_to"), m["type"], m.get("delay"), m.get("content"),
                      m.get("image_code"), m.get("image_mode"), m.get("replace_date")])
    sheet.freeze_panes = "A2"
    for col, width in {"A": 9, "B": 15, "C": 10, "D": 14, "E": 13, "F": 72, "G": 42, "H": 16, "I": 14}.items():
        sheet.column_dimensions[col].width = width
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
