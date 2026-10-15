import io
import unittest

from openpyxl import load_workbook

from app.core import export_xlsx
from app.version_archive import _markdown


class FictionalExportTests(unittest.TestCase):
    def test_usage_label_is_visible_without_changing_table_contract_or_old_exports(self):
        messages = [{'seq': 1, 'sender': '坐席', 'type': '文本', 'content': '你好'}]
        for config, title in (({}, '生成剧本'), ({'usage_context': 'fictional_rehearsal'}, '虚构排练剧本')):
            book = load_workbook(io.BytesIO(export_xlsx({'config': config, 'messages': messages})))
            self.assertEqual(book.active.title, title)
            self.assertEqual((book.active.max_row, book.active.max_column), (2, 9))
            self.assertEqual(book.active['F2'].value, '你好')
            if config:
                self.assertIn('虚构示例', book.properties.title)
                self.assertIn('不代表真实参与记录', book.properties.subject)
            book.close()
        record = {'version': 'R000', 'capture_status': 'running'}
        self.assertNotIn('虚构示例'.encode(), _markdown(record, messages))
        self.assertIn('不代表真实参与记录'.encode(), _markdown(record, messages, fictional=True))
