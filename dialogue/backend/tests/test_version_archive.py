from copy import deepcopy
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from app import version_archive
from app.version_archive import archive_candidate


class VersionArchiveTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.data = Path(self.temporary.name)
        self.messages = [{"seq": index, "sender": "坐席" if index == 1 else "客户1",
                          "content": f"第 {index} 句话", "reply_to": None,
                          "delay": "8", "stage": "自然群聊", "type": "文本",
                          "image_code": None, "image_mode": None, "replace_date": None}
                         for index in range(1, 151)]
        self.run = {"id": "abcdef0123456789abcdef0123456789", "status": "running",
                    "config": {"target_messages": 150, "dialogue_model": "gpt-6-sol"},
                    "messages": [{"content": "原运行中的消息保留"}],
                    "person_situations": {"客户1": {"personal_now": "正在等公交"}},
                    "metadata": {"mother": "母本中文.xlsx"},
                    "version_record": {"version": "V03", "code": "74b2429", "parent": "V02"}}

    def tearDown(self):
        self.temporary.cleanup()

    def capture(self, messages=None, reason="完整初稿"):
        return archive_candidate(self.data, self.run, self.messages if messages is None else messages, reason)

    def folder(self, record):
        return self.data / record["path"]

    def test_complete_snapshot_preserves_all_metadata_without_replacing_working_messages(self):
        original = deepcopy(self.run)
        record = self.capture()
        snapshot = json.loads((self.folder(record) / "run.json").read_text(encoding="utf-8"))
        self.assertEqual(snapshot["messages"], self.messages)
        self.assertEqual(self.run["messages"], original["messages"])
        for key in ("config", "person_situations", "metadata", "version_record"):
            self.assertEqual(snapshot[key], original[key])
            self.assertEqual(self.run[key], original[key])
        self.assertEqual(record["number"], 0)
        self.assertEqual(record["version"], "R000")
        self.assertEqual(record["code"], "74b2429")
        self.assertEqual(record["path"], f"versions/{self.run['id']}/R000-{record['hash'][:8]}")
        self.assertEqual(snapshot["candidate_versions"], [record])
        self.assertEqual(self.run["candidate_versions"], [record])

    def test_same_wording_keeps_first_status_reason_config_and_files(self):
        first = self.capture()
        before = {p.name: p.read_bytes() for p in self.folder(first).iterdir()}
        self.run["status"] = "completed"
        self.run["version_record"]["code"] = "bbbbbbb"
        self.run["config"]["dialogue_model"] = "different-model"
        second = self.capture(reason="最终完成")
        self.assertEqual(first, second)
        self.assertEqual(second["capture_status"], "running")
        self.assertEqual(second["reason"], "完整初稿")
        self.assertEqual(second["version_record"]["code"], "74b2429")
        self.assertEqual(len(self.run["candidate_versions"]), 1)
        self.assertEqual(before, {p.name: p.read_bytes() for p in self.folder(first).iterdir()})

    def test_each_export_relevant_field_changes_hash(self):
        variants = {"seq": 151, "sender": "客户2", "content": "不同台词", "reply_to": 1,
                    "delay": "12", "stage": "了解群用途", "type": "图片+文本",
                    "image_code": "IMG_NEW", "image_mode": "素材库", "replace_date": "2026-10-15"}
        records = [self.capture()]
        for field, value in variants.items():
            messages = deepcopy(self.messages)
            messages[-1][field] = value
            records.append(self.capture(messages, reason=field))
        self.assertEqual(len({record["hash"] for record in records}), 11)
        self.assertEqual([record["number"] for record in records], list(range(11)))
        self.assertEqual(len(self.run["candidate_versions"]), 11)

    def test_non_dialogue_message_metadata_does_not_make_duplicate(self):
        first = self.capture()
        messages = deepcopy(self.messages)
        messages[0]["source_rows"] = [42]
        messages[0]["debug_note"] = "not dialogue"
        self.assertEqual(self.capture(messages), first)

    def test_incomplete_draft_has_no_files_or_run_mutations(self):
        original = deepcopy(self.run)
        self.assertIsNone(self.capture(self.messages[:149]))
        self.assertEqual(self.run, original)
        self.assertEqual(list(self.data.iterdir()), [])

    def test_unsafe_run_ids_cannot_escape_data_directory(self):
        for run_id in ("../outside", "/tmp/outside", "..\\outside", "a" * 31, "g" * 32,
                       "a" * 32 + "/..", "A" * 32, "", None):
            with self.subTest(run_id=run_id):
                self.run["id"] = run_id
                with self.assertRaises(ValueError):
                    self.capture()
                self.assertNotIn("candidate_versions", self.run)
        self.assertEqual(list(self.data.iterdir()), [])

    def test_resolved_run_directory_cannot_escape_versions_root(self):
        resolve = Path.resolve
        outside = self.data / "outside"
        run_id = self.run["id"]
        def redirected(path, *args, **kwargs):
            if path.name == run_id:
                return outside
            return resolve(path, *args, **kwargs)
        with patch.object(Path, "resolve", redirected):
            with self.assertRaisesRegex(ValueError, "路径超出"):
                self.capture()
        self.assertFalse(outside.exists())

    def test_failed_write_keeps_older_draft_and_retry_reuses_next_number(self):
        first = self.capture()
        old_files = {p.name: p.read_bytes() for p in self.folder(first).iterdir()}
        old_index = deepcopy(self.run["candidate_versions"])
        messages = deepcopy(self.messages)
        messages[-1]["content"] = "准备修订的新台词"
        write = version_archive._write
        def fail_markdown(path, payload):
            if path.name == "messages.md":
                raise OSError("injected disk failure")
            write(path, payload)
        with patch.object(version_archive, "_write", fail_markdown):
            with self.assertRaises(OSError):
                self.capture(messages)
        self.assertEqual(self.run["candidate_versions"], old_index)
        self.assertEqual(old_files, {p.name: p.read_bytes() for p in self.folder(first).iterdir()})
        directory = self.folder(first).parent
        self.assertFalse(any(".tmp-" in path.name for path in directory.iterdir()))
        second = self.capture(messages)
        self.assertEqual(second["version"], "R001")
        self.assertEqual(len(self.run["candidate_versions"]), 2)

    def test_export_failure_never_publishes_or_appends_candidate_index(self):
        original = deepcopy(self.run)
        with patch.object(version_archive, "export_xlsx", side_effect=OSError("export failed")):
            with self.assertRaises(OSError):
                self.capture()
        self.assertEqual(self.run, original)
        self.assertFalse(any(self.data.rglob("manifest.json")))
        self.assertEqual(self.capture()["version"], "R000")

    def test_published_directory_recovers_lost_database_index_without_overwriting_status(self):
        first = self.capture()
        self.run.pop("candidate_versions")
        self.run["status"] = "completed"
        recovered = self.capture(reason="数据库持久化前中断后的重试")
        self.assertEqual(recovered, first)
        self.assertEqual(self.run["candidate_versions"], [first])
        snapshot = json.loads((self.folder(first) / "run.json").read_text(encoding="utf-8"))
        self.assertEqual(snapshot["status"], "running")

    def test_unindexed_older_archive_is_recovered_before_new_revision(self):
        first = self.capture()
        self.run.pop("candidate_versions")
        messages = deepcopy(self.messages)
        messages[-1]["content"] = "另一份完整稿"
        second = self.capture(messages)
        self.assertEqual(second["version"], "R001")
        self.assertEqual(self.run["candidate_versions"], [first, second])

    def test_image_placeholders_preserve_material_codes_and_blank_image_body(self):
        messages = deepcopy(self.messages)
        messages[3].update(type="图片", content="", image_code="SCREENSHOT_004", image_mode="待补素材")
        messages[4].update(type="图片+文本", content="这张图就是我看到的", image_code="SCREENSHOT_005")
        record = self.capture(messages)
        markdown = (self.folder(record) / "messages.md").read_text(encoding="utf-8")
        self.assertIn("[待补素材码：SCREENSHOT_004]", markdown)
        self.assertIn("[待补素材码：SCREENSHOT_005]", markdown)
        self.assertIn("这张图就是我看到的", markdown)

    def test_real_xlsx_has_151_rows_and_nine_columns_and_manifest_checksums(self):
        record = self.capture()
        folder = self.folder(record)
        book = load_workbook(io.BytesIO((folder / "messages.xlsx").read_bytes()))
        sheet = book.active
        self.assertEqual((sheet.max_row, sheet.max_column), (151, 9))
        self.assertEqual(sheet.cell(151, 6).value, "第 150 句话")
        manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["message_count"], 150)
        self.assertEqual(manifest["record"], record)
        for name, expected in manifest["files"].items():
            self.assertEqual(hashlib.sha256((folder / name).read_bytes()).hexdigest(), expected)

    def test_complete_structurally_invalid_draft_is_kept_without_invented_fields(self):
        messages = deepcopy(self.messages)
        del messages[-1]["type"]
        self.run["status"] = "failed"
        self.run["errors"] = ["最后一条消息缺少类型"]
        record = self.capture(messages, reason="结构检查失败的完整重写稿")
        snapshot = json.loads((self.folder(record) / "run.json").read_text(encoding="utf-8"))
        self.assertNotIn("type", snapshot["messages"][-1])
        self.assertEqual(snapshot["status"], "failed")
        self.assertEqual(snapshot["errors"], self.run["errors"])
        self.assertEqual(record["capture_status"], "failed")


if __name__ == "__main__":
    unittest.main()
